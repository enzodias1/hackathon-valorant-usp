"""Generates para_equipe/nomes_dos_times.xlsx: one row per team slug, with an automatic display-name
suggestion and the info a teammate needs to correct it (competitions, maps, most frequent nicks).

The filled sheet is read back by pipeline/teams.py (team display names + short tags for charts).
"""
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
P = ROOT / "data" / "processed"
OUT = ROOT / "para_equipe" / "nomes_dos_times.xlsx"

TOKENS = {
    "uff": "UFF", "ufmt": "UFMT", "ufu": "UFU", "uscs": "USCS", "fei": "FEI", "ime": "IME", "ceub": "CEUB", "caap": "CAAP",
    "ufpr": "UFPR", "ufg": "UFG", "pucgo": "PUC-GO", "umc": "UMC", "usp": "USP", "utfpr": "UTFPR", "a2e": "A2E", "bbn": "BBN",
    "ufmg": "UFMG", "unifesp": "UNIFESP", "fatec": "FATEC", "pg": "PG", "ibmec": "Ibmec", "ufrj": "UFRJ", "gc": "GC",
    "unirio": "UNIRIO", "ufpb": "UFPB", "gdu": "GDU", "iff": "IFF", "pucc": "PUC-Campinas", "maua": "Mauá", "unicamp": "Unicamp",
    "fametro": "FAMETRO", "inatel": "Inatel", "poli": "Poli", "unievangelica": "UniEVANGÉLICA", "fenix": "Fênix",
    "panterao": "Panterão", "erex": "EREX", "b": "B",
}


def suggest(slug):
    return " ".join(TOKENS.get(t, t.capitalize()) for t in slug.split("_"))


def main():
    pm = pd.read_parquet(P / "player_match.parquet")
    m = pd.read_parquet(P / "matches.parquet")
    pm = pm.merge(m[["match_id", "competition", "stage"]], on="match_id")
    pm["comp"] = pm.competition + pm.stage.fillna("").map(lambda s: f" ({s})" if s else "")
    g = pm.groupby("team")
    df = pd.DataFrame({
        "slug (não editar)": g.size().index,
        "nome sugerido (automático)": [suggest(s) for s in g.size().index],
        "NOME CORRETO (preencher só se o sugerido estiver errado)": "",
        "SIGLA CURTA p/ gráficos (até 8 letras)": "",
        "INSTITUIÇÃO (nome completo)": "",
        "mapas jogados": g.match_id.nunique().values,
        "competições": g.comp.apply(lambda s: "; ".join(sorted(set(s)))).values,
        "nicks mais frequentes (p/ identificar o time)": g.name.apply(lambda s: ", ".join(s.str.split("#").str[0].value_counts().head(6).index)).values,
        "observação": "",
    }).sort_values("mapas jogados", ascending=False)

    OUT.parent.mkdir(exist_ok=True)
    with pd.ExcelWriter(OUT, engine="openpyxl") as xw:
        df.to_excel(xw, sheet_name="Times", index=False)
        ws = xw.sheets["Times"]
        widths = [26, 30, 34, 22, 36, 10, 60, 60, 30]
        fill_in = PatternFill("solid", fgColor="FFF2CC")
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = w
        for c in ws[1]:
            c.font = Font(bold=True)
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for col in (3, 4, 5, 9):
            for row in range(1, ws.max_row + 1):
                ws.cell(row, col).fill = fill_in
        ws.row_dimensions[1].height = 45
        ws.freeze_panes = "B2"
    print(f"{OUT} ({len(df)} times)")


if __name__ == "__main__":
    main()
